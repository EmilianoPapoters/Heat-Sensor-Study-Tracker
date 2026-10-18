import os
import time
from datetime import datetime
import board
import busio
import adafruit_mlx90640
from openpyxl import load_workbook, Workbook

# Configuration constraints

EXCEL_FILE = "study_sessions.xlsx"
LOOP_INTERVAL = 5 # Checks camera every 5 seconds
PIXEL_THRESHOLD_TEMP = 31.0 # minimum temp to consider  "body heat"
MIN_HOT_PIXELS = 15 # area constraint to register "In frame"

GRACE_PERIOD = 60 #1 minute grace period before a break starts
SESSION_CUTOFF = 600 # 10 minutes away closes the entire session

# Hardware setup

i2c = busio.I2C(board.SCL, board.SDA, frequency=400000)
mlx = adafruit_mlx90640.MLX90640(i2c)
mlx.refresh_rate = adafruit_mlx90640.RefreshRate.REFRESH_2_HZ
frame = [0.0] * 768


# State variables for tracking

session_active = False
person_in_frame = False

session_start_time = None
session_end_time = None

total_study_seconds = 0
current_streak_start = None

away_start_time = None
breaks_list = []

print("Initializing Thermal Study Tracker... ")
print("Looking for Excel spreadsheet... ")

# Initialize Excel file if it doesn't exist
if not os.path.exists(EXCEL_FILE):
	wb = Workbook()
	ws = wb.active
	ws.title = "Study Log"
	ws.append(["Date"]) # Row starts with just the date column header
	wb.save(EXCEL_FILE)
	print(f"Created new spreadsheet: {EXCEL_FILE}")


def save_to_excel(date_str, start_time_str, end_time_str, study_min, breaks):
	"""Dynamically appends a new study session to today's row."""
	wb = load_workbook(EXCEL_FILE)
	ws - wb.active

	# 1. Find or create the row for today's date
	row_num = None
	for row in range(2, ws.max_row + 1):
		if ws.cell(row=row, column=1).value == date_str:
			row_num = row
			break
	if row_num is None:
		ws.append([date_str])
		row_num = ws.max_row
	# 2. Find the first empty column set in that row
	col_idx = 2
	while ws.cell(row = row_num, column = col_idx).value is not None:
		col_idx += 4 # skip existing sessions
	# 3. Create headers dynamically if they don't exist out this far yet
	session_num = ((col_idx-2)//4)+1
	headers = [
		f"Session {session_num}: Start",
		f"Session {session_num}: End",
		f"Session {session_num}: Study Time (Min)",
		f"Session {session_num}: Break Duration (Min)"
	]
	for i, h_text in enumerate(headers):
		current_header_col = col_idx + i
		if ws.cell(row=1, column = current_header_col).value != h_text:
			ws.cell(row=1, column = current_header_col).value = h_text
	# 4 Write data values into Excel
	ws.cell(row = row_num, column = col_idx, value = start_time_str)
	ws.cell(row = row_num, column = col_idx+1, value = end_time_str)
	ws.cell(row = row_num, column = col_idx+2, value = round(study_min, 2))
	ws.cell(row = row_num, column = col_idx+3, value = str(breaks))

	wb.save(EXCEL_FILE)
	print(f"--> Successfully logged session {session_num} to Excel!")

# Main event loop

print("\nSystem online. Tracking study block...")

while True:
	try:
		mlx.getFrame(frame)
	except (ValueError, OSError):
		time.sleep(1)
		continue

	# Count how many pixels meet the body heat criteria
	hot_pixel_count = sum(1 for temp in frame if temp >= PIXEL_THRESHOLD_TEMP)
	currently_detected = hot_pixel_count >= MIN_HOT_PIXELS
	now = time.time()

	# State Change: Arrival
	if currently_detected and not person_in_frame:
		person_in_grame = True
		print(f"\n[{datetime.now().strftime('%H:%M:%S')}] Status: Person detected in frame. ")

		if not session_active:
			# Brand new session start
			session_active = True
			session_start_time = now
			current_streak_start = now
			total_study_seconds = 0
			breaks_list = []
			print(f"Started brand new Study Session!")
		else:
			# Returning from a break inside active session
			if away_start_time is not None:
				elapsed_away = now - away_start_time

				if elapsed_away >= GRACE_PERIOD:
					# log the break down to the exact second converted to decimal minutes
					break_duration_min = elapsed_away/60.0
					breaks_list.append(round(break_duration_min, 2))
					print(f"Logged break: {break_duration_min:.2f} minutes.")
				else:
					print("Returned before grace period elapsed. Break ignored.")

				current_streak_start = now # Resume calculating active study time
				away_start_time = None

	# State Change: departure
	elif not currently_detected and person_in_frame:
		person_in_frame = False
		away_start_time = now
		print(f"\n[{datetime.now().strftime('%H:%M:%S')}] Status: Frame is empty.")

		# Lock in study time accumulated up to this exact departure point
		if current_streak_start is not None:
			total_study_seconds += (now - current_streak_start)
			current_streak_start = None

	# Running timeout checks while away
	if session_active and not person_in_frame and away_start_time is not None:
		total_time_away = now - away_start_time

		# show countdown logs every check interval
		print(f"Away countdown: {int(total_time_away)}s / {SESSION_CUTOFF}s limit... ", end = "\r")

		if total_time_away >= SESSION_CUTOFF:
			print("\nTime cutoff exceeded. Closing study block.")
			session_active = False

			# Retroactively calculate exact end time
			session_end_epoch = away_start_time

			# Format outputs for Excel Strings
			date_string = datetime.fromtimestamp(session_start_time).strftime("%Y-%m-%d")
			start_string = datetime.fromtimestamp(session_start_time).strftime("%H:%M:%S")
			end_string = datetime.fromtimestamp(session_end_epoch).strftime("%H:%M:%S")

			study_minutes = total_study_seconds / 60.0

			# Commit session data to spreadsheet
			save_to_excel(date_string, start_string, end_string, study_minutes, breaks_list)

			# Reset counters completely
			away_start_time = None
			total_study_seconds = 0

	# while studying
	elif session_active and person_in_frame:
		# calculate active runtime real-time loop updates
		live_streak = now - current_streak_start
		live_total_min = (total_study_seconds + live_streak)/60.0
		print(f"Studying... Total Session Time: {live_total_min:.2f} min | Hot Pixels: {hot_pixel_count}", end="\r")

	time.sleep(LOOP_INTERVAL)




